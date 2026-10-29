#!/usr/bin/env python3
"""Exercise a prepared work-tools tree in place, offline, and check it stays byte-exact.

Usage: python3 scripts/work-tools-smoke.py build/work-tools/<native target>
Runs from the read-only tree with a PATH without node, npm, pnpm, bun, python or
git, and without PYTHONPYCACHEPREFIX: a missing .pyc would land in the tree and
fail the final inventory check. Launches no Cua process and no browser.
"""

import importlib.util
import json
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import tempfile


spec = importlib.util.spec_from_file_location("work_tools", Path(__file__).with_name("work-tools.py"))
tools = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tools)


DOCUMENT_PROBE = """
import json,sys,importlib.metadata
from pathlib import Path
from openpyxl import Workbook,load_workbook
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
import pdfplumber
from PIL import Image
from reportlab.pdfgen.canvas import Canvas
root=Path(sys.argv[1])
book=Workbook();book.active['A1']='Terminus';book.save(root/'sample.xlsx')
assert load_workbook(root/'sample.xlsx').active['A1'].value=='Terminus'
doc=Document();doc.add_paragraph('Terminus');doc.save(root/'sample.docx')
assert Document(root/'sample.docx').paragraphs[0].text=='Terminus'
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[0]);slide.shapes.title.text='Terminus';slides.save(root/'sample.pptx')
assert Presentation(root/'sample.pptx').slides[0].shapes.title.text=='Terminus'
pdf=Canvas(str(root/'sample.pdf'));pdf.drawString(72,720,'Terminus');pdf.save()
assert 'Terminus' in PdfReader(root/'sample.pdf').pages[0].extract_text()
with pdfplumber.open(root/'sample.pdf') as opened: assert 'Terminus' in opened.pages[0].extract_text()
Image.new('RGB',(16,16),'red').save(root/'sample.png')
with Image.open(root/'sample.png') as image: assert image.size==(16,16)
print(json.dumps({'python':sys.version.split()[0],'round_trips':['xlsx','docx','pptx','pdf','png'],'distributions':{d.metadata['Name']:d.version for d in importlib.metadata.distributions()}}))
"""

# Lo que el usuario tenga instalado no puede hacer pasar la prueba.
HIDDEN = ("node", "nodejs", "npm", "npx", "pnpm", "bun", "bunx", "python", "pip", "git", "uv", "rg")


def native_target():
    machine = {"x86_64": "x86_64", "AMD64": "x86_64", "arm64": "aarch64", "aarch64": "aarch64"}[platform.machine()]
    system = {"Linux": "unknown-linux-gnu", "Darwin": "apple-darwin", "Windows": "pc-windows-msvc"}[platform.system()]
    return f"{machine}-{system}"


def bare_path(lab):
    """A PATH of the system's own tools, without any language runtime, git or search tool."""
    if os.name == "nt":
        return str(Path(os.environ["SystemRoot"]) / "System32")
    bin = lab / "system-bin"
    bin.mkdir()
    for directory in ["/usr/bin", "/bin", "/usr/sbin", "/sbin"]:
        for tool in sorted(Path(directory).iterdir()) if Path(directory).is_dir() else []:
            if tool.name.startswith(HIDDEN) or (bin / tool.name).exists():
                continue
            (bin / tool.name).symlink_to(tool)
    return str(bin)


def exercise(directory, target, lab, bare):
    """Run every bundled tool from directory/<tree> and return what each one answered."""
    manifest = tools.load_manifest(directory, target)
    tree = directory / manifest["tree"]
    entries = {entry["id"]: tree / entry["id"] / entry["executable"]
               for entry in manifest["entries"] if entry["executable"]}
    home = lab / "home"
    home.mkdir(exist_ok=True)
    env = {key: os.environ[key] for key in ("SystemRoot", "COMSPEC", "PATHEXT") if key in os.environ}
    env.update(HOME=str(home), USERPROFILE=str(home), PATH=bare, LANG="C.UTF-8", TMPDIR=str(lab), TMP=str(lab),
               TEMP=str(lab), UV_CACHE_DIR=str(lab / "uv-cache"), UV_PYTHON_DOWNLOADS="never", UV_OFFLINE="1",
               GIT_CONFIG_NOSYSTEM="1", GIT_CONFIG_GLOBAL=os.devnull)

    def run(command, cwd=lab, extra=None):
        result = subprocess.run([str(arg) for arg in command], cwd=cwd, env=dict(env, **(extra or {})), text=True,
                                stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
        if result.returncode:
            raise RuntimeError(f"{Path(str(command[0])).name} failed: {result.stderr[-2000:]}")
        return result.stdout.strip()

    for hidden in ("node", "npm", "pnpm", "bun", "python3", "git"):
        if shutil.which(hidden, path=bare):
            raise RuntimeError(f"The bare PATH still offers {hidden}")
    versions = {id: run([path, "--version"]).splitlines()[0] for id, path in entries.items()}
    for entry in manifest["entries"]:
        if entry["executable"] and entry["version"].split("+")[0] not in versions[entry["id"]]:
            raise RuntimeError(f"{entry['id']} answers {versions[entry['id']]}, pinned {entry['version']}")
    git = entries["git"]
    if os.name != "nt":
        # Como el motor: el Git de dugite no sabe dónde vive sin estas dos variables.
        env.update(GIT_EXEC_PATH=str(tree / "git/libexec/git-core"),
                   GIT_TEMPLATE_DIR=str(tree / "git/share/git-core/templates"))
    repo = lab / "repo"
    run([git, "init", repo])
    (repo / "file.txt").write_text("Terminus offline\n")
    run([git, "add", "file.txt"], repo)
    run([git, "-c", "user.name=Terminus smoke", "-c", "user.email=smoke@example.invalid", "commit", "-m", "smoke"], repo)
    # Clonar por file:// lanza git-upload-pack desde el exec-path incluido.
    run([git, "clone", repo.as_uri(), lab / "clone"])
    if (lab / "clone" / "file.txt").read_text() != "Terminus offline\n":
        raise RuntimeError("Bundled Git did not clone over file://")
    if "Terminus offline" not in run([entries["ripgrep"], "Terminus", repo]):
        raise RuntimeError("Bundled ripgrep did not find the text")
    python_dir = tree / "python"
    aliases = ([python_dir / "python.exe", python_dir / "python3.exe"] if os.name == "nt"
               else [python_dir / "bin/python", python_dir / "bin/python3"])
    reports = []
    for index, alias in enumerate(aliases):
        output = lab / f"documents-{index}"
        output.mkdir()
        reports.append(json.loads(run([alias, "-I", "-c", DOCUMENT_PROBE, output])))
    distributions = reports[0]["distributions"]
    for package in manifest["documents"]["packages"]:
        normal = lambda value: re.sub(r"[-_.]+", "-", value).lower()
        name = next((n for n in distributions if normal(n) == normal(package["name"])), None)
        if name is None or distributions[name] != package["version"]:
            raise RuntimeError(f"Bundled Python lacks {package['name']}=={package['version']}")
    blocked = subprocess.run([str(aliases[0]), "-I", "-m", "pip", "install", "--no-index", "six"], env=env,
                             capture_output=True, text=True, timeout=120)
    if blocked.returncode == 0 or "externally-managed" not in (blocked.stdout + blocked.stderr):
        raise RuntimeError("The bundled Python accepted a package install")
    run([entries["uv"], "venv", "--python", aliases[0], lab / "extra"])
    run([entries["agent-browser"], "--help"])
    report = tools.verify_output(directory, target)
    return {"target": target, "tree": manifest["tree"], "versions": versions, "git": "commit and file:// clone",
            "ripgrep": "search", "documents": reports[0]["round_trips"] + ["pdfplumber"],
            "python_aliases": [a.name for a in aliases], "pep668": "blocked", "uv": "venv from the bundled Python",
            "inventory_after_use": "identical", "files": report["files"]}


def main():
    if len(sys.argv) != 2:
        raise ValueError(__doc__)
    directory = Path(sys.argv[1]).resolve()
    target = native_target()
    if directory.name != target:
        raise ValueError(f"This smoke runs the native target {target}")
    with tempfile.TemporaryDirectory(prefix="terminus-tools-smoke-") as temp:
        lab = Path(temp)
        print(json.dumps(exercise(directory, target, lab, bare_path(lab)), indent=2))


if __name__ == "__main__":
    try:
        main()
    except (ValueError, OSError, subprocess.SubprocessError, RuntimeError) as error:
        print(f"work-tools-smoke: {error}", file=sys.stderr)
        sys.exit(1)
